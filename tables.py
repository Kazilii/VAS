from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
import re

wb = Workbook()
ws = wb.active

inventoryfile = []

with open('inventory.txt') as f:
    for line in f:
        line = line.strip('\n')
        inventoryfile.append(line)

banqinv = []
for item in inventoryfile:
    banqinv.append(["{}".format(item), "", "", ""])

# Create column headings. MUST BE STRINGS
ws.append(['Item Name', 'Previous', 'Quantity', 'Cost'])
for row in banqinv:
    ws.append(row)

tab = Table(displayName="Table1", ref="A1:D1")

# Add a default style with striped rows and banded columns

style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=True,
                       showLastColumn=True, showRowStripes=True, showColumnStripes=True)
tab.tableStyleInfo = style
ws.add_table(tab)

 # Used for the next for statement
def as_text(value):
    if value is None:
        return ""
    return str(value)

 # Automatically determine the width for any column based on the longest cell
for column_cells in ws.columns:
    length = max(len(as_text(cell.value)) for cell in column_cells)
    ws.column_dimensions[column_cells[0].column].width = length + (length * 0.2)

def rownumber(row):
    regex = r".(\d+)"

    matches = re.finditer(regex, row)

    for matchNum, match in enumerate(matches):
        matchNum = matchNum + 1

    for groupNum in range(0, len(match.groups())):
        groupNum = groupNum + 1

    rownum = match.group(groupNum)

    return(rownum)

for row in ws.rows: #Run the next bit of code on every row in the sheet
    for cell in row: # Run the next bit of code on every cell in the row
        if "Tuaca" in cell.value: # Check if the test string is in the cell
            print("Success!")
            rownum = rownumber(str(row)) #Get the number of the row
            cellk = ws.cell(row=int(rownum), column=3) #Select the cell we want to edit, in this case the "Quantity" field
            cellk.value = "1" # Set the value
            print(rownumber(str(row)))

wb.save('BANQINV.xlsx')